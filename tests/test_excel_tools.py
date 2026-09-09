"""
Unit tests for the Excel arm (excel_tools.py), fully offline -- openpyxl
reads/writes real .xlsx files on disk, no browser or LLM involved.
"""
import openpyxl
import pytest

from excel_tools import ExcelError, ExcelSession, ExcelToolProvider


def _make_workbook(path, sheet_data):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    for row in sheet_data:
        ws.append(row)
    wb.save(path)


def test_open_read_write_save_roundtrip(tmp_path):
    path = tmp_path / "data.xlsx"
    _make_workbook(path, [["Name", "Price"], ["Widget", 10], ["Gadget", 20]])

    session = ExcelSession()
    assert session.is_open() is False

    result = session.execute("excel_open", {"path": str(path)})
    assert "Sheet1" in result
    assert session.is_open() is True

    result = session.execute("excel_read_cell", {"sheet": "Sheet1", "cell": "A2"})
    assert "Widget" in result

    result = session.execute("excel_read_range", {"sheet": "Sheet1", "cell_range": "A1:B3"})
    assert "Widget" in result and "Gadget" in result

    session.execute("excel_write_cell", {"sheet": "Sheet1", "cell": "B2", "value": "15"})
    result = session.execute("excel_read_cell", {"sheet": "Sheet1", "cell": "B2"})
    assert "15" in result

    session.execute("excel_save", {})

    # Re-open fresh to prove the write actually landed on disk, not just in memory.
    reopened = openpyxl.load_workbook(path)
    assert reopened["Sheet1"]["B2"].value == 15
    assert isinstance(reopened["Sheet1"]["B2"].value, int)  # numeric coercion, not the string "15"


def test_create_if_missing(tmp_path):
    path = tmp_path / "new.xlsx"
    session = ExcelSession()

    with pytest.raises(ExcelError):
        session.execute("excel_open", {"path": str(path)})

    result = session.execute("excel_open", {"path": str(path), "create_if_missing": True})
    assert "did not exist" in result
    session.execute("excel_save", {})
    assert path.exists()


def test_unknown_sheet_and_missing_workbook_raise_clear_errors(tmp_path):
    path = tmp_path / "data.xlsx"
    _make_workbook(path, [["A"]])
    session = ExcelSession()

    with pytest.raises(ExcelError, match="No workbook is open"):
        session.execute("excel_read_cell", {"sheet": "Sheet1", "cell": "A1"})

    session.execute("excel_open", {"path": str(path)})
    with pytest.raises(ExcelError, match="No sheet named"):
        session.execute("excel_read_cell", {"sheet": "DoesNotExist", "cell": "A1"})


def test_list_sheets(tmp_path):
    path = tmp_path / "multi.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "First"
    wb.create_sheet("Second")
    wb.save(path)

    session = ExcelSession()
    session.execute("excel_open", {"path": str(path)})
    result = session.execute("excel_list_sheets", {})
    assert "First" in result and "Second" in result


def test_write_cell_coerces_numeric_strings_but_leaves_text_as_text(tmp_path):
    path = tmp_path / "coerce.xlsx"
    session = ExcelSession()
    session.execute("excel_open", {"path": str(path), "create_if_missing": True})

    session.execute("excel_write_cell", {"sheet": "Sheet", "cell": "A1", "value": "3.5"})
    session.execute("excel_write_cell", {"sheet": "Sheet", "cell": "A2", "value": "42"})
    session.execute("excel_write_cell", {"sheet": "Sheet", "cell": "A3", "value": "Widget"})
    session.execute("excel_save", {})

    wb = openpyxl.load_workbook(path)
    assert wb["Sheet"]["A1"].value == 3.5 and isinstance(wb["Sheet"]["A1"].value, float)
    assert wb["Sheet"]["A2"].value == 42 and isinstance(wb["Sheet"]["A2"].value, int)
    assert wb["Sheet"]["A3"].value == "Widget"


def test_excel_save_with_an_explicit_path_saves_there_instead_of_overwriting(tmp_path):
    original = tmp_path / "original.xlsx"
    target = tmp_path / "copy.xlsx"
    session = ExcelSession()
    session.execute("excel_open", {"path": str(original), "create_if_missing": True})
    session.execute("excel_write_cell", {"sheet": "Sheet", "cell": "A1", "value": "hello"})

    result = session.execute("excel_save", {"path": str(target)})

    assert str(target) in result
    assert target.exists()
    assert not original.exists()  # never written -- excel_save went to `target` instead
    assert session.path == target  # subsequent excel_save (no path) would now target `target`


def test_excel_save_creates_missing_parent_directories(tmp_path):
    nested = tmp_path / "a" / "b" / "c" / "data.xlsx"
    session = ExcelSession()
    session.execute("excel_open", {"path": str(nested), "create_if_missing": True})
    session.execute("excel_save", {})
    assert nested.exists()


def test_opening_a_corrupted_file_raises_a_clear_excel_error(tmp_path):
    path = tmp_path / "not_really_xlsx.xlsx"
    path.write_text("this is not a valid xlsx file")
    session = ExcelSession()

    with pytest.raises(ExcelError, match="Could not open"):
        session.execute("excel_open", {"path": str(path)})


def test_close_resets_the_session_to_unopened(tmp_path):
    path = tmp_path / "data.xlsx"
    session = ExcelSession()
    session.execute("excel_open", {"path": str(path), "create_if_missing": True})
    assert session.is_open() is True

    session.close()

    assert session.is_open() is False
    assert session.path is None
    with pytest.raises(ExcelError, match="No workbook is open"):
        session.execute("excel_list_sheets", {})


def test_describe_for_confirmation_names_the_explicit_path_when_given(tmp_path):
    session = ExcelSession()
    session.execute("excel_open", {"path": str(tmp_path / "data.xlsx"), "create_if_missing": True})
    provider = ExcelToolProvider(session)

    other_path = str(tmp_path / "elsewhere.xlsx")
    desc = provider.describe_for_confirmation("excel_save", {"path": other_path})

    assert other_path in desc
    assert str(session.path) not in desc


def test_describe_for_confirmation_names_the_current_file_when_no_path_given(tmp_path):
    path = tmp_path / "data.xlsx"
    session = ExcelSession()
    session.execute("excel_open", {"path": str(path), "create_if_missing": True})
    provider = ExcelToolProvider(session)

    desc = provider.describe_for_confirmation("excel_save", {})

    assert str(path) in desc
